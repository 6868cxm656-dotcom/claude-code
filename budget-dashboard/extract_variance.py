"""Parse a Xero 'Income and Expenditure Detailed Budget Variance Report' export
into JSON for the dashboard.

Expected layout (Xero default): title rows 1-3, header row 5 with
Account / Budget - Full year / Reforecast - Full Year / Reforecast - N Months /
Actual N months / Variance / Variance %. Outline levels: column A = section,
column B = group, column C = '#### - Account name' lines.
"""
import json, re, sys
import openpyxl

SRC = sys.argv[1]
OUT = sys.argv[2]

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb[wb.sheetnames[0]]

section, group = None, None
out = []
for row in ws.iter_rows(min_row=6):
    a, b, c = row[0].value, row[1].value, row[2].value
    if a and not str(a).startswith("Total"):
        section = str(a).strip()
        group = None
    if b and not str(b).startswith("Total"):
        group = str(b).strip()
    if c and re.match(r"^\d{4}", str(c)):
        name = str(c).strip()
        vals = [row[i].value if isinstance(row[i].value, (int, float)) else 0
                for i in range(3, 9)]
        out.append({
            "section": section, "group": group,
            "code": int(name[:4]),
            "name": re.sub(r"^\d{4}\s*-\s*", "", name),
            "fy_budget": round(vals[0], 2), "fy_rf": round(vals[1], 2),
            "ytd_rf": round(vals[2], 2), "ytd_act": round(vals[3], 2),
            "var": round(vals[4], 2), "var_pct": round(vals[5], 4),
        })

with open(OUT, "w") as f:
    json.dump(out, f)
print(f"{len(out)} lines -> {OUT}")
